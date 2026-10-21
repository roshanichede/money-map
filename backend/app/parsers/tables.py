"""Turn any spreadsheet-like statement into rows of text: CSV, Excel (.xlsx / .xls) or the HTML tables some
banks save with an .xls name. Indian banks in particular export Excel with account details above the table."""
import csv
import io
from datetime import date, datetime
from html.parser import HTMLParser

from .common import ParseError


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.2f}"
    return str(value).replace("\xa0", " ").strip()


def _keep(rows: list[list[str]]) -> list[list[str]]:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        raise ParseError("The file is empty.")
    return rows


def read_csv(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4000]
    delimiter = ","
    if sample.count("\t") > sample.count(","):
        delimiter = "\t"
    elif sample.count(";") > sample.count(",") * 2:
        delimiter = ";"
    return _keep([[c.strip() for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)])


def read_xlsx(content: bytes) -> list[list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    return _keep([[_text(v) for v in row] for row in ws.iter_rows(values_only=True)])


def read_xls(content: bytes) -> list[list[str]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        out = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                out.append(xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
            else:
                out.append(_text(cell.value))
        rows.append(out)
    return _keep(rows)


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def read_html(content: bytes) -> list[list[str]]:
    p = _Tables()
    p.feed(content.decode("utf-8", errors="replace"))
    return _keep(p.rows)


def kind_of(filename: str, content: bytes) -> str:
    name = filename.lower()
    head = content[:512].lstrip()
    if content[:5] == b"%PDF-" or name.endswith(".pdf"):
        return "pdf"
    if head[:1] == b"<" or b"<table" in content[:4000].lower():
        return "html"
    if content[:2] == b"PK":
        return "xlsx"
    if content[:4] == b"\xd0\xcf\x11\xe0":
        return "xls"
    if name.endswith((".xlsx", ".xls")):
        raise ParseError("That Excel file couldn't be opened. Try saving it again as .xlsx or CSV.")
    return "csv"


READERS = {"csv": read_csv, "xlsx": read_xlsx, "xls": read_xls, "html": read_html}
